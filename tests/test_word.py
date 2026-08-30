from pathlib import Path

import pytest
from docx import Document


def document_text(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.extend(cell.text for cell in row.cells)
    return "\n".join(blocks)


def test_render_stops_on_blocking_markdown_issues(
    plugin_root: Path, fixtures: Path, syllabus_module, tmp_path: Path
):
    output = tmp_path / "invalid.docx"
    template = (
        plugin_root
        / "skills"
        / "course-syllabus-compiler"
        / "assets"
        / "syllabus-master.docx"
    )

    with pytest.raises(ValueError, match="blocking"):
        syllabus_module.render_syllabus(
            fixtures / "invalid-syllabus.md",
            fixtures / "valid-assessment.md",
            template,
            output,
        )

    assert not output.exists()


def test_one_master_renders_two_distinct_courses(
    plugin_root: Path, fixtures: Path, syllabus_module, tmp_path: Path
):
    template = (
        plugin_root
        / "skills"
        / "course-syllabus-compiler"
        / "assets"
        / "syllabus-master.docx"
    )
    output_one = tmp_path / "one.docx"
    output_two = tmp_path / "two.docx"

    result_one = syllabus_module.render_syllabus(
        fixtures / "valid-syllabus.md",
        fixtures / "valid-assessment.md",
        template,
        output_one,
    )
    result_two = syllabus_module.render_syllabus(
        fixtures / "valid-syllabus-2.md",
        fixtures / "valid-assessment-2.md",
        template,
        output_two,
    )

    assert result_one["inspection"].clean
    assert result_two["inspection"].clean
    text_one = document_text(output_one)
    text_two = document_text(output_two)
    assert "测试课程一" in text_one and "测试课程二" not in text_one
    assert "测试课程二" in text_two and "测试课程一" not in text_two
    assert "课程评价" in text_one
    assert "课程评价" in text_two


def test_render_rejects_stale_assessment_materialization(
    plugin_root: Path, fixtures: Path, syllabus_module, tmp_path: Path
):
    template = (
        plugin_root
        / "skills"
        / "course-syllabus-compiler"
        / "assets"
        / "syllabus-master.docx"
    )
    stale = tmp_path / "stale-assessment.md"
    stale.write_text(
        (fixtures / "valid-assessment.md")
        .read_text(encoding="utf-8")
        .replace("test-assessment-001", "stale-hash"),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="assessment"):
        syllabus_module.render_syllabus(
            fixtures / "valid-syllabus.md",
            stale,
            template,
            tmp_path / "stale.docx",
        )


def test_render_reuses_unchanged_build_state(
    plugin_root: Path, fixtures: Path, syllabus_module, tmp_path: Path
):
    template = (
        plugin_root
        / "skills"
        / "course-syllabus-compiler"
        / "assets"
        / "syllabus-master.docx"
    )
    output = tmp_path / "cached.docx"
    state = tmp_path / "build-state.json"

    first = syllabus_module.render_syllabus(
        fixtures / "valid-syllabus.md",
        fixtures / "valid-assessment.md",
        template,
        output,
        state_path=state,
    )
    second = syllabus_module.render_syllabus(
        fixtures / "valid-syllabus.md",
        fixtures / "valid-assessment.md",
        template,
        output,
        state_path=state,
    )

    assert not first["reused"]
    assert second["reused"]


def test_inspection_rejects_ai_authorship_marker(syllabus_module, tmp_path: Path):
    path = tmp_path / "marked.docx"
    document = Document()
    document.add_paragraph("本大纲由AI生成")
    document.save(path)

    result = syllabus_module.inspect_docx(path)

    assert not result.clean
    assert "AI_AUTHORSHIP_MARKER" in result.codes


def test_inspection_allows_ai_as_course_subject(syllabus_module, tmp_path: Path):
    path = tmp_path / "course-content.docx"
    document = Document()
    document.add_paragraph("人工智能导论讲授AI算法基础。")
    document.save(path)

    result = syllabus_module.inspect_docx(path)

    assert result.clean
    assert "AI_AUTHORSHIP_MARKER" not in result.codes
