import json
from pathlib import Path
import shutil
import subprocess

from docx import Document
import pytest


def run_cache(script: Path, source: Path, cache: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-File",
            str(script),
            "-SourceRoot",
            str(source),
            "-CacheRoot",
            str(cache),
            "-Json",
        ],
        text=True,
        encoding="utf-8",
        capture_output=True,
    )


def test_unchanged_markdown_is_reused_and_each_run_is_logged(
    tmp_path: Path, plugin_root: Path, fixtures: Path
):
    source = tmp_path / "source"
    source.mkdir()
    shutil.copy2(fixtures / "source-course.md", source / "course.md")
    cache = tmp_path / "cache"
    script = (
        plugin_root
        / "skills"
        / "course-foundation-builder"
        / "scripts"
        / "update_markdown_cache.ps1"
    )

    first = run_cache(script, source, cache)
    second = run_cache(script, source, cache)

    assert first.returncode == 0, first.stderr
    assert second.returncode == 0, second.stderr
    assert json.loads(first.stdout)["converted"] == 1
    assert json.loads(second.stdout)["reused"] == 1

    manifest = json.loads((cache / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["files"][0]["route"] == "direct-markdown"
    assert manifest["files"][0]["tool"] == "existing-markdown"
    assert (cache / manifest["files"][0]["cache_path"]).exists()

    log_lines = (cache / "conversion-log.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(log_lines) == 2
    assert json.loads(log_lines[1])["reused"] == 1


@pytest.mark.skipif(shutil.which("markitdown") is None, reason="MarkItDown is not installed")
def test_docx_is_converted_with_markitdown(tmp_path: Path, plugin_root: Path):
    source = tmp_path / "source"
    source.mkdir()
    document = Document()
    document.add_heading("Reusable course source", level=1)
    document.add_paragraph("A deterministic conversion fixture.")
    document.save(source / "course.docx")
    cache = tmp_path / "cache"
    script = (
        plugin_root
        / "skills"
        / "course-foundation-builder"
        / "scripts"
        / "update_markdown_cache.ps1"
    )

    result = run_cache(script, source, cache)

    assert result.returncode == 0, result.stderr
    manifest = json.loads((cache / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["files"][0]["route"] == "markitdown-docx"
    cached = cache / manifest["files"][0]["cache_path"]
    assert "Reusable course source" in cached.read_text(encoding="utf-8")


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_html_is_converted_with_pandoc(tmp_path: Path, plugin_root: Path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "course.html").write_text(
        "<h1>Course evidence</h1><p>Reusable HTML fixture.</p>", encoding="utf-8"
    )
    cache = tmp_path / "cache"
    script = (
        plugin_root
        / "skills"
        / "course-foundation-builder"
        / "scripts"
        / "update_markdown_cache.ps1"
    )

    result = run_cache(script, source, cache)

    assert result.returncode == 0, result.stderr
    manifest = json.loads((cache / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["files"][0]["route"] == "pandoc-html"
    cached = cache / manifest["files"][0]["cache_path"]
    assert "# Course evidence" in cached.read_text(encoding="utf-8")
