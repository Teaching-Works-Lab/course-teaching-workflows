from pathlib import Path
import importlib.util
import subprocess
import sys
import zipfile
from xml.etree import ElementTree as ET

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "syllabus.py"


def load_module(name: str, rel: str):
    path = ROOT / rel
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(module)
    return module


def test_renderer_creates_school_basic_info_and_dynamic_sections(tmp_path):
    renderer = load_module("render_school_syllabus", "scripts/render_school_syllabus.py")
    source = ROOT / "examples" / "25JD31401-artificial-intelligence.md"
    out = tmp_path / "ai.docx"
    renderer.render_syllabus(source, out)
    doc = Document(out)
    assert len(doc.tables) >= 4
    assert len(doc.tables[0].rows) == 12
    assert len(doc.tables[0].columns) == 8
    first_table_text = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "课程简介" in first_table_text
    assert "总学时" in first_table_text and "理论" in first_table_text and "上机" in first_table_text
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "理论教学环节及基本要求" in body
    assert "上机教学环节及基本要求" in body


def test_renderer_preserves_textbook_terminal_isbn(tmp_path):
    renderer = load_module("render_school_syllabus", "scripts/render_school_syllabus.py")
    source = ROOT / "examples" / "25JD32408-intelligent-equipment.md"
    out = tmp_path / "equipment.docx"
    renderer.render_syllabus(source, out)
    doc = Document(out)
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    idx = paragraphs.index("（一）选用教材")
    assert paragraphs[idx + 1].endswith("9787111776741")
    assert not paragraphs[idx + 1].endswith(("。", "."))


def test_docx_uses_valid_twip_grid_widths_and_is_clean(tmp_path):
    renderer = load_module("render_school_syllabus", "scripts/render_school_syllabus.py")
    inspector = load_module("inspect_docx", "scripts/inspect_docx.py")
    source = ROOT / "examples" / "25JD31401-artificial-intelligence.md"
    out = tmp_path / "ai.docx"
    renderer.render_syllabus(source, out)
    report = inspector.inspect_docx(out)
    assert report["issues"] == [], report
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(out) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    widths = [int(e.attrib[f"{{{ns['w']}}}w"]) for e in root.findall(".//w:tblGrid/w:gridCol", ns)]
    assert widths
    assert max(widths) < 30000


def test_syllabus_cli_check_render_inspect(tmp_path):
    source = ROOT / "examples" / "25JD31401-artificial-intelligence.md"
    out = tmp_path / "cli.docx"
    commands = [
        [sys.executable, str(SCRIPT), "check", "--input", str(source)],
        [sys.executable, str(SCRIPT), "render", "--input", str(source), "--output", str(out)],
        [sys.executable, str(SCRIPT), "inspect", "--input", str(out)],
    ]
    for command in commands:
        result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True)
        assert result.returncode == 0, result.stdout + result.stderr
    assert out.exists()


def test_renderer_writes_exactly_one_page_number_field(tmp_path):
    renderer = load_module("render_school_syllabus", "scripts/render_school_syllabus.py")
    source = ROOT / "examples" / "25JD31401-artificial-intelligence.md"
    out = tmp_path / "page-number.docx"
    renderer.render_syllabus(source, out)
    with zipfile.ZipFile(out) as zf:
        footer_xml = "\n".join(
            zf.read(name).decode("utf-8", errors="replace")
            for name in zf.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )
    assert footer_xml.count("w:fldSimple") == 1
