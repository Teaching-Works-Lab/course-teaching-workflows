from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
from typing import Sequence
import zipfile
from xml.etree import ElementTree as ET

from docx import Document

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"
COLLEGES = {
    "经济管理学院",
    "机械与动力工程学院",
    "土木建筑工程学院",
    "电子与电气工程学院",
    "计算机与信息技术学院",
    "轨道交通学院",
    "艺术学院",
    "外国语学院",
    "化学与制药工程学院",
}


def _visible_text(document: Document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def inspect_docx(path: Path) -> dict[str, object]:
    path = Path(path)
    issues: list[dict[str, str]] = []

    def add(code: str, message: str) -> None:
        issues.append({"code": code, "message": message})

    try:
        document = Document(path)
    except Exception as exc:  # pragma: no cover - defensive CLI path
        return {"path": str(path), "issues": [{"code": "DOCX_OPEN", "message": str(exc)}]}

    if not document.tables:
        add("NO_TABLES", "Word文档没有表格")
    else:
        first = document.tables[0]
        if len(first.rows) != 12 or len(first.columns) != 8:
            add("BASIC_INFO_DIMENSIONS", f"课程基本信息表应为12×8，实际为{len(first.rows)}×{len(first.columns)}")
        first_text = "\n".join(cell.text for row in first.rows for cell in row.cells)
        for required in ("课程名称（中英文）", "课程学时及分配", "课程简介", "大纲更新时间"):
            if required not in first_text:
                add("BASIC_INFO_SCHEMA", f"课程基本信息表缺少：{required}")

    section = document.sections[0]
    if abs(section.page_width.cm - 21.0) > 0.1 or abs(section.page_height.cm - 29.7) > 0.1:
        add("PAGE_GEOMETRY", "页面不是A4纵向")

    visible = _visible_text(document)
    if "计信学院" in visible:
        add("COLLEGE_ABBREVIATION", "Word中不得使用简称“计信学院”")
    if "授课学院" in visible and not any(name in visible for name in COLLEGES):
        add("COLLEGE_NOT_IN_WHITELIST", "授课学院未使用学校正式全称")

    paragraphs = [p.text.strip() for p in document.paragraphs]
    try:
        index = paragraphs.index("（一）选用教材")
        textbook = paragraphs[index + 1]
        if not re.search(r"ISBN\s+\d{13}$", textbook):
            add("TEXTBOOK_TERMINAL_ISBN", "Word选用教材必须以ISBN数字结束")
    except (ValueError, IndexError):
        add("TEXTBOOK_MISSING", "Word缺少选用教材段落")

    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        xml_bytes = zf.read("word/document.xml")
        xml_text = xml_bytes.decode("utf-8", errors="replace")
        if any(name.startswith("word/comments") for name in names) or "<w:commentReference" in xml_text:
            add("COMMENTS_PRESENT", "Word终稿包含批注")
        if re.search(r"<w:(?:ins|del|moveFrom|moveTo)(?:\s|>)", xml_text):
            add("TRACKED_CHANGES_PRESENT", "Word终稿包含修订")
        if "<w:vanish" in xml_text or "<w:webHidden" in xml_text:
            add("HIDDEN_TEXT_PRESENT", "Word终稿包含隐藏文字")
        if any(marker in visible for marker in ("AI生成", "由AI生成", "ChatGPT生成", "Codex生成", "source_id", "to_confirm")):
            add("GENERATION_MARKER", "Word终稿包含生成或内部标记")
        if "Noto Serif CJK" in xml_text or "Noto Sans CJK" in xml_text:
            add("NON_WINDOWS_FONT", "最终Word写入了Linux专用Noto CJK字体")

        root = ET.fromstring(xml_bytes)
        widths: list[int] = []
        for element in root.findall(".//w:tblGrid/w:gridCol", {"w": W_NS}):
            raw = element.attrib.get(W + "w")
            if raw and raw.isdigit():
                widths.append(int(raw))
        if widths and max(widths) >= 30000:
            add("GRIDCOL_WIDTH_SUSPICIOUS", f"表格列宽疑似EMU/twip混用：最大值{max(widths)}")

        metadata = ""
        for name in ("docProps/core.xml", "docProps/custom.xml", "docProps/app.xml"):
            if name in names:
                metadata += zf.read(name).decode("utf-8", errors="replace")
        if re.search(r"OpenAI|ChatGPT|Codex|GPT-[0-9]", metadata, re.IGNORECASE):
            add("TOOL_METADATA", "Word属性包含生成工具元数据")

    return {"path": str(path), "issues": issues}


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Inspect a final syllabus DOCX")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    report = inspect_docx(args.input)
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    elif report["issues"]:
        for issue in report["issues"]:  # type: ignore[index]
            print(f"{issue['code']}: {issue['message']}")
    else:
        print("PASS")
    return 2 if report["issues"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
