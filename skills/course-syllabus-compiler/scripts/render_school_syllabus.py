from __future__ import annotations

import argparse
from pathlib import Path
import re
from typing import Iterable, Sequence, NamedTuple

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SCHOOL_GRID = [1931, 1634, 1215, 1104, 477, 627, 1033, 1033]
CONTENT_WIDTH = sum(SCHOOL_GRID)
BODY_CN = "宋体"
HEADING_CN = "黑体"
LATIN = "Times New Roman"


class TableBlock(NamedTuple):
    headers: list[str]
    rows: list[list[str]]


class Section(NamedTuple):
    title: str
    body: str


def _clean_title(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip())


def parse_document(text: str) -> tuple[str, list[Section]]:
    title_match = re.search(r"^#\s+(.+?)\s*$", text, re.MULTILINE)
    if not title_match:
        raise ValueError("Markdown must contain a level-1 syllabus title")
    title = _clean_title(title_match.group(1))
    matches = list(re.finditer(r"^##\s+(.+?)\s*$", text, re.MULTILINE))
    if not matches:
        raise ValueError("Markdown must contain level-2 syllabus sections")
    sections: list[Section] = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections.append(Section(_clean_title(match.group(1)), text[start:end].strip()))
    return title, sections


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _separator(line: str) -> bool:
    values = _cells(line)
    return bool(values) and all(re.fullmatch(r":?-{3,}:?", value.replace(" ", "")) for value in values)


def parse_table(lines: Sequence[str], start: int) -> tuple[TableBlock, int]:
    if start + 1 >= len(lines) or not lines[start].lstrip().startswith("|") or not _separator(lines[start + 1]):
        raise ValueError("Not a Markdown table")
    headers = _cells(lines[start])
    rows: list[list[str]] = []
    index = start + 2
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        values = _cells(lines[index])
        if len(values) != len(headers):
            raise ValueError(f"Markdown table has inconsistent column count near line {index + 1}")
        rows.append(values)
        index += 1
    return TableBlock(headers, rows), index


def first_table(body: str) -> TableBlock:
    lines = body.splitlines()
    for index, line in enumerate(lines):
        if line.lstrip().startswith("|") and index + 1 < len(lines) and _separator(lines[index + 1]):
            table, _ = parse_table(lines, index)
            return table
    raise ValueError("Expected a Markdown table")


def _strip_markdown(text: str) -> str:
    text = text.replace("<br>", "\n")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    return text.strip()


def parse_basic_info(body: str) -> tuple[dict[str, str], dict[str, str], str]:
    table = first_table(body)
    fields: dict[str, str] = {}
    for row in table.rows:
        if len(row) >= 2:
            fields[_strip_markdown(row[0])] = _strip_markdown(row[1])

    intro: dict[str, str] = {}
    intro_match = re.search(r"^###\s+课程简介\s*$([\s\S]*)", body, re.MULTILINE)
    intro_text = intro_match.group(1) if intro_match else ""
    labels = ["课程基本定位", "核心学习结果", "主要教学方法"]
    for index, label in enumerate(labels):
        pattern = re.compile(rf"^\*\*{re.escape(label)}\*\*\s*$", re.MULTILINE)
        match = pattern.search(intro_text)
        if not match:
            continue
        end = len(intro_text)
        for next_label in labels[index + 1 :]:
            next_match = re.search(rf"^\*\*{re.escape(next_label)}\*\*\s*$", intro_text[match.end() :], re.MULTILINE)
            if next_match:
                end = match.end() + next_match.start()
                break
        update_match = re.search(r"^\*\*大纲更新时间\*\*", intro_text[match.end() :], re.MULTILINE)
        if update_match:
            end = min(end, match.end() + update_match.start())
        intro[label] = re.sub(r"\s*\n\s*", "", intro_text[match.end() : end].strip())
    update = ""
    update_match = re.search(r"\*\*大纲更新时间\*\*\s*[：:]\s*([^\n]+)", intro_text)
    if update_match:
        update = update_match.group(1).strip()
    return fields, intro, update


def parse_hours(value: str) -> dict[str, str]:
    result = {name: "0" for name in ["总学时", "理论", "实践", "实验", "上机", "项目式"]}
    aliases = {"项目": "项目式", "理论学时": "理论", "实践学时": "实践", "实验学时": "实验", "上机学时": "上机"}
    for part in re.split(r"[；;，,]", value):
        part = part.strip()
        match = re.match(r"([^0-9]+?)\s*([0-9]+(?:\.[0-9]+)?)\s*(?:学时)?$", part)
        if not match:
            continue
        key = match.group(1).strip().replace("：", "").replace(":", "")
        key = aliases.get(key, key)
        if key in result:
            result[key] = match.group(2)
    return result


def _set_run_font(run, size: float, *, bold: bool = False, chinese: str = BODY_CN) -> None:
    run.font.name = LATIN
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr, value in (("ascii", LATIN), ("hAnsi", LATIN), ("eastAsia", chinese)):
        rfonts.set(qn(f"w:{attr}"), value)


def _configure_paragraph(paragraph, *, align=None, line_spacing: float = 1.18, before: float = 0, after: float = 0) -> None:
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)


def _add_inline(paragraph, text: str, *, size: float, default_bold: bool = False, chinese: str = BODY_CN) -> None:
    cursor = 0
    for match in re.finditer(r"\*\*(.*?)\*\*|`([^`]*)`", text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            _set_run_font(run, size, bold=default_bold, chinese=chinese)
        content = match.group(1) if match.group(1) is not None else match.group(2)
        run = paragraph.add_run(content)
        _set_run_font(run, size, bold=(match.group(1) is not None) or default_bold, chinese=chinese)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        _set_run_font(run, size, bold=default_bold, chinese=chinese)
    if not paragraph.runs:
        run = paragraph.add_run("")
        _set_run_font(run, size, bold=default_bold, chinese=chinese)


def _clear_cell(cell) -> None:
    cell.text = ""
    if not cell.paragraphs:
        cell.add_paragraph()


def _set_cell_markdown(cell, text: str, *, size: float = 9, bold: bool = False, center: bool = False) -> None:
    _clear_cell(cell)
    segments = text.replace("<br />", "<br>").split("<br>")
    for index, segment in enumerate(segments):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        _configure_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.12)
        _add_inline(paragraph, segment.strip(), size=size, default_bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_cell_margins(cell, top: int = 80, start: int = 70, bottom: int = 80, end: int = 70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def _set_table_grid(table, widths: Sequence[int]) -> None:
    if len(widths) != len(table.columns):
        raise ValueError(f"Grid width count {len(widths)} does not match table columns {len(table.columns)}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    grid = table._tbl.tblGrid
    existing = list(grid)
    for child in existing:
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width)))
        grid.append(col)

    for row in table.rows:
        seen: set[int] = set()
        column = 0
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            tc_pr = cell._tc.get_or_add_tcPr()
            span_node = tc_pr.find(qn("w:gridSpan"))
            span = int(span_node.get(qn("w:val"))) if span_node is not None else 1
            width = sum(widths[column : column + span])
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            _set_cell_margins(cell)
            column += span
            if column >= len(widths):
                break


def _new_table(document: DocumentType, rows: int, cols: int, widths: Sequence[int]):
    table = document.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)
    _set_table_grid(table, widths)
    return table


def _set_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    # A master can already contain a PAGE field. Clear all existing footer
    # content before adding the single canonical field, otherwise WPS/Word
    # can display duplicated values such as 11, 22, 33.
    p_element = paragraph._p
    for child in list(p_element):
        if child.tag != qn("w:pPr"):
            p_element.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    _set_run_font(run, 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def _configure_document(document: DocumentType) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)
    _set_page_number(section)
    normal = document.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_CN)


def _empty_master(path: Path) -> None:
    doc = Document()
    _configure_document(doc)
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "课程教学大纲空白主文档"
    doc.save(path)


def create_master(path: Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    _empty_master(path)


def _document_from_master(master: Path | None) -> DocumentType:
    document = Document(master) if master and Path(master).exists() else Document()
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    _configure_document(document)
    return document


def _add_title(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    _configure_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    run = paragraph.add_run(text)
    _set_run_font(run, 16, bold=True, chinese=HEADING_CN)


def _add_level1(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    _configure_paragraph(paragraph, before=5, after=3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    _set_run_font(run, 14, bold=True, chinese=HEADING_CN)


def _add_level2(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    _configure_paragraph(paragraph, before=4, after=2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    _set_run_font(run, 12, bold=True)


def _add_minor_heading(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    _configure_paragraph(paragraph, before=3, after=1)
    paragraph.paragraph_format.keep_with_next = True
    _add_inline(paragraph, text, size=10.5, default_bold=True)


def _add_body(document: DocumentType, text: str, *, indent: bool = False, center: bool = False) -> None:
    paragraph = document.add_paragraph()
    _configure_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT, after=1)
    if indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    _add_inline(paragraph, text, size=10.5)


def _merge_range(row, start: int, end: int):
    return row.cells[start].merge(row.cells[end])


def render_basic_info(document: DocumentType, fields: dict[str, str], intro: dict[str, str], update: str) -> None:
    table = _new_table(document, 12, 8, SCHOOL_GRID)

    row = table.rows[0]
    _set_cell_markdown(row.cells[0], "课程名称（中英文）", bold=True, center=True)
    _set_cell_markdown(_merge_range(row, 1, 7), fields.get("课程名称（中英文）", ""), center=True, size=9.5)

    pairs = [
        ("课程代码", fields.get("课程代码", ""), "课程类别", fields.get("课程类别", "")),
        ("课程性质", fields.get("课程性质", ""), "授课语言", fields.get("授课语言", "")),
        ("授课学期", fields.get("授课学期", fields.get("开课学期", "")), "学分", fields.get("学分", "")),
    ]
    for row_index, (left_label, left_value, right_label, right_value) in enumerate(pairs, start=1):
        row = table.rows[row_index]
        _set_cell_markdown(row.cells[0], left_label, bold=True, center=True)
        _set_cell_markdown(_merge_range(row, 1, 2), left_value, center=True)
        _set_cell_markdown(row.cells[3], right_label, bold=True, center=True)
        _set_cell_markdown(_merge_range(row, 4, 7), right_value, center=True)

    row4 = table.rows[4]
    row5 = table.rows[5]
    label_top = _merge_range(row4, 0, 1)
    label_bottom = _merge_range(row5, 0, 1)
    label = label_top.merge(label_bottom)
    _set_cell_markdown(label, "课程学时及分配", bold=True, center=True)
    hour_names = ["总学时", "理论", "实践", "实验", "上机", "项目式"]
    hours = parse_hours(fields.get("课程学时及分配", ""))
    for offset, name in enumerate(hour_names, start=2):
        _set_cell_markdown(row4.cells[offset], name, bold=True, center=True, size=8.8)
        _set_cell_markdown(row5.cells[offset], hours.get(name, "0"), center=True)

    for row_index, label_text, value in [
        (6, "适用专业", fields.get("适用专业", "")),
        (7, "授课学院", fields.get("授课学院", fields.get("开课单位", ""))),
        (8, "先修课程", fields.get("先修课程", "无")),
        (9, "后续课程", fields.get("后续课程", "无")),
    ]:
        row = table.rows[row_index]
        _set_cell_markdown(row.cells[0], label_text, bold=True, center=True)
        _set_cell_markdown(_merge_range(row, 1, 7), value, size=9.2)

    row = table.rows[10]
    _set_cell_markdown(row.cells[0], "课程简介", bold=True, center=True)
    intro_cell = _merge_range(row, 1, 7)
    _clear_cell(intro_cell)
    for index, label_text in enumerate(["课程基本定位", "核心学习结果", "主要教学方法"]):
        paragraph = intro_cell.paragraphs[0] if index == 0 else intro_cell.add_paragraph()
        _configure_paragraph(paragraph, after=2, line_spacing=1.16)
        heading_run = paragraph.add_run(label_text + "\n")
        _set_run_font(heading_run, 9.4, bold=True)
        body_run = paragraph.add_run(intro.get(label_text, ""))
        _set_run_font(body_run, 9.2)
    intro_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    row = table.rows[11]
    _set_cell_markdown(row.cells[0], "大纲更新时间", bold=True, center=True)
    _set_cell_markdown(_merge_range(row, 1, 7), update, center=True)

    _set_table_grid(table, SCHOOL_GRID)
    for row in table.rows:
        _set_cant_split(row)


def _scaled_widths(weights: Sequence[float], total: int = CONTENT_WIDTH) -> list[int]:
    raw = [total * value / sum(weights) for value in weights]
    widths = [max(350, int(round(value))) for value in raw]
    widths[-1] += total - sum(widths)
    return widths


def table_widths(headers: Sequence[str]) -> list[int]:
    count = len(headers)
    normalized = [re.sub(r"\s+", "", header) for header in headers]
    if count == 3 and "课程总目标" in normalized[0]:
        return [1050, 4650, CONTENT_WIDTH - 5700]
    if count == 5 and any("学时" in header for header in normalized):
        return [1300, 1900, 600, 4000, CONTENT_WIDTH - 7800]
    if count == 4:
        return [1700, 900, 5100, CONTENT_WIDTH - 7700]
    if count == 7 and any("90" in header for header in normalized):
        return _scaled_widths([1.1, 0.7, 1.45, 1.45, 1.45, 1.45, 1.45])
    if count >= 8:
        weights = [1.45] + [0.7 if "学时" in header else 0.9 for header in normalized[1:-1]] + [1.35]
        return _scaled_widths(weights)
    if count == 2:
        return [2100, CONTENT_WIDTH - 2100]
    return _scaled_widths([1.0] * count)


def render_table(document: DocumentType, block: TableBlock, *, font_size: float | None = None) -> None:
    widths = table_widths(block.headers)
    table = _new_table(document, 1, len(block.headers), widths)
    size = font_size if font_size is not None else (8.1 if len(block.headers) >= 8 else 8.8)
    for index, header in enumerate(block.headers):
        _set_cell_markdown(table.rows[0].cells[index], header, bold=True, center=True, size=size)
    _set_repeat_header(table.rows[0])
    _set_cant_split(table.rows[0])
    for values in block.rows:
        row = table.add_row()
        for index, value in enumerate(values):
            center = index > 0 and ("学时" in block.headers[index] or "权重" in block.headers[index] or "成绩占比" in block.headers[index])
            _set_cell_markdown(row.cells[index], value, size=size, center=center)
    _set_table_grid(table, widths)


def render_goals(document: DocumentType, body: str) -> None:
    block = first_table(body)
    render_table(document, block, font_size=8.8)


def _text_after_first_table(body: str) -> str:
    lines = body.splitlines()
    for index, line in enumerate(lines):
        if line.lstrip().startswith("|") and index + 1 < len(lines) and _separator(lines[index + 1]):
            _, end = parse_table(lines, index)
            return "\n".join(lines[end:]).strip()
    return body.strip()


def render_teaching(document: DocumentType, body: str) -> None:
    block = first_table(body)
    render_table(document, block, font_size=8.5)
    tail = _text_after_first_table(body)
    for line in tail.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("**") and stripped.endswith("**"):
            _add_minor_heading(document, _strip_markdown(stripped))
        else:
            _add_body(document, _strip_markdown(stripped))


def parse_blocks(body: str) -> list[tuple[str, object]]:
    lines = body.splitlines()
    blocks: list[tuple[str, object]] = []
    index = 0
    paragraph_lines: list[str] = []

    def flush() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            text = " ".join(line.strip() for line in paragraph_lines if line.strip()).strip()
            if text:
                blocks.append(("paragraph", text))
            paragraph_lines = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and _separator(lines[index + 1]):
            flush()
            table, index = parse_table(lines, index)
            blocks.append(("table", table))
            continue
        if stripped.startswith("#### "):
            flush()
            blocks.append(("h4", stripped[5:].strip()))
        elif stripped.startswith("### "):
            flush()
            blocks.append(("h3", stripped[4:].strip()))
        elif stripped.startswith("**") and stripped.endswith("**") and len(stripped) > 4:
            flush()
            blocks.append(("minor", _strip_markdown(stripped)))
        elif re.match(r"^(?:\d+\.|\[\d+\])\s+", stripped):
            flush()
            blocks.append(("list", stripped))
        elif stripped == "---":
            flush()
        else:
            paragraph_lines.append(stripped)
        index += 1
    flush()
    return blocks


def render_generic(document: DocumentType, body: str) -> None:
    for kind, payload in parse_blocks(body):
        if kind == "h3":
            _add_level2(document, str(payload))
        elif kind == "h4":
            _add_minor_heading(document, str(payload))
        elif kind == "minor":
            _add_minor_heading(document, str(payload))
        elif kind == "table":
            render_table(document, payload)  # type: ignore[arg-type]
        elif kind == "list":
            _add_body(document, _strip_markdown(str(payload)))
        elif kind == "paragraph":
            text = _strip_markdown(str(payload))
            center = bool(re.fullmatch(r"\d{4}年\d{1,2}月", text))
            _add_body(document, text, indent=not center, center=center)


def render_syllabus(source: Path, output: Path, master: Path | None = None) -> Path:
    source = Path(source)
    output = Path(output)
    title, sections = parse_document(source.read_text(encoding="utf-8"))
    if master is None:
        candidate = Path(__file__).resolve().parents[1] / "assets" / "syllabus-master.docx"
        master = candidate if candidate.exists() else None
    document = _document_from_master(master)
    _add_title(document, title)

    for section in sections:
        _add_level1(document, section.title)
        if "课程基本信息" in section.title:
            fields, intro, update = parse_basic_info(section.body)
            render_basic_info(document, fields, intro, update)
        elif "课程目标" in section.title:
            render_goals(document, section.body)
        elif "教学环节及基本要求" in section.title:
            render_teaching(document, section.body)
        else:
            render_generic(document, section.body)

    props = document.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.keywords = ""
    props.subject = ""
    props.title = title
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Render school-format course syllabus DOCX")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--master", type=Path)
    parser.add_argument("--create-master", type=Path)
    args = parser.parse_args(argv)
    if args.create_master:
        create_master(args.create_master)
        print(args.create_master)
        return 0
    render_syllabus(args.input, args.output, args.master)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
