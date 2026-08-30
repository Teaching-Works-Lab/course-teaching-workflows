from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from typing import Mapping
import zipfile
from xml.etree import ElementTree

import yaml
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SCRIPT_VERSION = "0.2.0"
RULE_VERSION = "1"
REQUIRED_FIELDS = (
    "课程中文名称",
    "课程代码",
    "学分",
    "总学时",
    "适用专业",
    "开课单位",
)
HOUR_FIELDS = ("理论学时", "实践学时", "实验学时", "上机学时", "项目学时")
VALID_LEVELS = {"basic", "changed", "enhanced"}
CHINESE_NUMERALS = "一二三四五六七八九十"
AI_AUTHORSHIP_PATTERNS = (
    re.compile(r"(?:本大纲|本课程大纲|本文档|本文件).{0,12}(?:由)?(?:AI|ChatGPT|Codex|模型).{0,8}(?:生成|编写|撰写)", re.I),
    re.compile(r"(?:AI|ChatGPT|Codex|模型)(?:自动)?(?:生成|编写|撰写)(?:的大纲|本文档|本文件)?", re.I),
)


@dataclass(frozen=True)
class Issue:
    code: str
    severity: str
    message: str
    location: str


@dataclass(frozen=True)
class InspectionResult:
    clean: bool
    codes: tuple[str, ...]
    findings: tuple[str, ...]
    visible_text: str


def _front_matter(text: str) -> tuple[dict[str, object], str]:
    if not text.startswith("---"):
        raise ValueError("Syllabus must start with YAML front matter")
    parts = text.split("---", 2)
    if len(parts) != 3:
        raise ValueError("YAML front matter is not closed")
    metadata = yaml.safe_load(parts[1]) or {}
    if not isinstance(metadata, dict):
        raise ValueError("YAML front matter must be a mapping")
    return metadata, parts[2]


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_separator(line: str) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell) for cell in _cells(line))


def _parse_tables(body: str) -> dict[str, list[dict[str, str]]]:
    lines = body.splitlines()
    tables: dict[str, list[dict[str, str]]] = {}
    heading = ""
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("#"):
            heading = line.lstrip("#").strip()
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and _is_separator(lines[index + 1]):
            headers = _cells(line)
            index += 2
            rows: list[dict[str, str]] = []
            while index < len(lines) and lines[index].startswith("|"):
                values = _cells(lines[index])
                if len(values) != len(headers):
                    raise ValueError(f"Table under '{heading}' has inconsistent columns")
                rows.append(dict(zip(headers, values, strict=True)))
                index += 1
            tables[heading] = rows
            continue
        index += 1
    return tables


def read_syllabus(path: Path) -> dict[str, object]:
    path = Path(path)
    text = path.read_text(encoding="utf-8")
    metadata, body = _front_matter(text)
    formal = body.split("# 正式大纲内容", 1)[1] if "# 正式大纲内容" in body else body
    formal = formal.split("# 内部追踪区", 1)[0]
    return {
        "path": path,
        "text": text,
        "metadata": metadata,
        "tables": _parse_tables(body),
        "formal_text": formal,
    }


def read_assessment(path: Path) -> dict[str, object]:
    path = Path(path)
    text = path.read_text(encoding="utf-8")
    metadata, body = _front_matter(text)
    return {
        "path": path,
        "text": text,
        "metadata": metadata,
        "tables": _parse_tables(body),
    }


def _section_prose(text: str, heading: str) -> str:
    pattern = re.compile(rf"^##\s+{re.escape(heading)}\s*$", re.MULTILINE)
    match = pattern.search(text)
    if not match:
        return ""
    start = match.end()
    next_heading = re.search(r"^#{1,2}\s+", text[start:], re.MULTILINE)
    end = start + next_heading.start() if next_heading else len(text)
    lines = []
    for line in text[start:end].strip().splitlines():
        if line.startswith("|"):
            break
        if line.strip():
            lines.append(line.strip())
    return "\n".join(lines)


def _table(data: Mapping[str, object], name: str) -> list[dict[str, str]]:
    tables = data.get("tables", {})
    if not isinstance(tables, dict):
        return []
    rows = tables.get(name, [])
    return rows if isinstance(rows, list) else []


def _integer(value: str) -> int | None:
    value = value.strip().removesuffix("%").strip()
    if not re.fullmatch(r"-?\d+", value):
        return None
    return int(value)


def _split_ids(value: str) -> set[str]:
    return {part for part in re.split(r"[,，、;；\s]+", value.strip()) if part}


def _duplicates(values: list[str]) -> set[str]:
    seen: set[str] = set()
    repeated: set[str] = set()
    for value in values:
        if not value:
            continue
        if value in seen:
            repeated.add(value)
        seen.add(value)
    return repeated


def _issue(code: str, severity: str, message: str, location: str) -> Issue:
    return Issue(code=code, severity=severity, message=message, location=location)


def check_syllabus(data: Mapping[str, object], level: str = "basic") -> list[Issue]:
    if level not in VALID_LEVELS:
        raise ValueError(f"Unsupported review level: {level}")

    issues: list[Issue] = []
    info_rows = _table(data, "课程基本信息")
    info = {row.get("字段", ""): row for row in info_rows}

    for field in REQUIRED_FIELDS:
        row = info.get(field, {})
        value = row.get("内容", "").strip()
        status = row.get("状态", "").strip()
        if not value or status == "to_confirm":
            issues.append(
                _issue(
                    "MISSING_REQUIRED_FIELD",
                    "blocking",
                    f"Required field is missing or unresolved: {field}",
                    f"课程基本信息/{field}",
                )
            )

    total = _integer(info.get("总学时", {}).get("内容", ""))
    components = [_integer(info.get(field, {}).get("内容", "0")) for field in HOUR_FIELDS]
    if total is not None and all(value is not None for value in components):
        component_total = sum(value for value in components if value is not None)
        if total != component_total:
            issues.append(
                _issue(
                    "HOURS_MISMATCH",
                    "blocking",
                    f"Total hours {total} do not equal component hours {component_total}",
                    "课程基本信息/学时",
                )
            )

    goals = _table(data, "课程目标")
    units = _table(data, "教学单元")
    assessments = _table(data, "课程评价")
    criteria = _table(data, "评分标准")

    goal_ids = {row.get("goal_id", "") for row in goals if row.get("goal_id", "")}
    unit_ids = {row.get("unit_id", "") for row in units if row.get("unit_id", "")}
    assessment_ids = {
        row.get("assessment_id", "") for row in assessments if row.get("assessment_id", "")
    }

    for name, values in (
        ("goal_id", [row.get("goal_id", "") for row in goals]),
        ("unit_id", [row.get("unit_id", "") for row in units]),
        ("assessment_id", [row.get("assessment_id", "") for row in assessments]),
        ("criterion_id", [row.get("criterion_id", "") for row in criteria]),
    ):
        for duplicate in sorted(_duplicates(values)):
            issues.append(
                _issue("DUPLICATE_ID", "blocking", f"Duplicate {name}: {duplicate}", name)
            )

    assessment_weights = [_integer(row.get("权重", "")) for row in assessments]
    if assessment_weights and all(value is not None for value in assessment_weights):
        weight_total = sum(value for value in assessment_weights if value is not None)
        if weight_total != 100:
            issues.append(
                _issue(
                    "ASSESSMENT_WEIGHT_MISMATCH",
                    "blocking",
                    f"Assessment weights total {weight_total}, expected 100",
                    "课程评价/权重",
                )
            )

    criteria_totals: dict[str, int] = {}
    for row in criteria:
        assessment_id = row.get("assessment_id", "")
        weight = _integer(row.get("指标权重", ""))
        if assessment_id and weight is not None:
            criteria_totals[assessment_id] = criteria_totals.get(assessment_id, 0) + weight
        if assessment_id and assessment_id not in assessment_ids:
            issues.append(
                _issue(
                    "UNKNOWN_ASSESSMENT_REFERENCE",
                    "blocking",
                    f"Criterion references unknown assessment: {assessment_id}",
                    "评分标准",
                )
            )
    for assessment_id in sorted(assessment_ids):
        if criteria_totals.get(assessment_id) != 100:
            issues.append(
                _issue(
                    "CRITERION_WEIGHT_MISMATCH",
                    "blocking",
                    f"Criteria for {assessment_id} total {criteria_totals.get(assessment_id, 0)}, expected 100",
                    f"评分标准/{assessment_id}",
                )
            )

    taught_goals: set[str] = set()
    for row in units:
        references = _split_ids(row.get("对应目标", ""))
        taught_goals.update(references)
        for reference in sorted(references - goal_ids):
            issues.append(
                _issue(
                    "UNKNOWN_GOAL_REFERENCE",
                    "blocking",
                    f"Teaching unit references unknown goal: {reference}",
                    f"教学单元/{row.get('unit_id', '')}",
                )
            )

    assessed_goals: set[str] = set()
    for row in assessments:
        goal_references = _split_ids(row.get("覆盖目标", ""))
        unit_references = _split_ids(row.get("覆盖单元", ""))
        assessed_goals.update(goal_references)
        for reference in sorted(goal_references - goal_ids):
            issues.append(
                _issue(
                    "UNKNOWN_GOAL_REFERENCE",
                    "blocking",
                    f"Assessment references unknown goal: {reference}",
                    f"课程评价/{row.get('assessment_id', '')}",
                )
            )
        for reference in sorted(unit_references - unit_ids):
            issues.append(
                _issue(
                    "UNKNOWN_UNIT_REFERENCE",
                    "blocking",
                    f"Assessment references unknown unit: {reference}",
                    f"课程评价/{row.get('assessment_id', '')}",
                )
            )

    for goal_id in sorted(goal_ids - taught_goals):
        issues.append(
            _issue(
                "GOAL_NOT_TAUGHT",
                "blocking",
                f"Course goal is not supported by a teaching unit: {goal_id}",
                f"课程目标/{goal_id}",
            )
        )
    for goal_id in sorted(goal_ids - assessed_goals):
        issues.append(
            _issue(
                "GOAL_NOT_ASSESSED",
                "blocking",
                f"Course goal is not covered by an assessment: {goal_id}",
                f"课程目标/{goal_id}",
            )
        )

    formal_text = str(data.get("formal_text", ""))
    if "[待填写]" in formal_text or "to_confirm" in formal_text:
        issues.append(
            _issue(
                "UNRESOLVED_PLACEHOLDER",
                "blocking",
                "Formal syllabus content contains unresolved placeholders",
                "正式大纲内容",
            )
        )

    metadata = data.get("metadata", {})
    if isinstance(metadata, dict):
        if not metadata.get("assessment_version") or not metadata.get("assessment_hash"):
            issues.append(
                _issue(
                    "ASSESSMENT_SYNC_MISSING",
                    "blocking",
                    "Assessment version or hash is missing",
                    "控制区",
                )
            )

    if level == "enhanced":
        generic_pattern = re.compile(r"^第(?:[一二三四五六七八九十]+|\d+)章$")
        for row in units:
            unit_name = row.get("单元名称", "").strip()
            if generic_pattern.fullmatch(unit_name):
                issues.append(
                    _issue(
                        "SEMANTIC_GENERIC_UNIT_NAME",
                        "warning",
                        f"Teaching unit name is too generic to communicate its content: {unit_name}",
                        f"教学单元/{row.get('unit_id', '')}",
                    )
                )

    return issues


def _signature(input_path: Path, level: str, config: Mapping[str, object]) -> str:
    digest = hashlib.sha256()
    digest.update(Path(input_path).read_bytes())
    digest.update(level.encode("utf-8"))
    digest.update(RULE_VERSION.encode("ascii"))
    digest.update(SCRIPT_VERSION.encode("ascii"))
    digest.update(json.dumps(config, ensure_ascii=False, sort_keys=True).encode("utf-8"))
    return digest.hexdigest()


def check_file(
    input_path: Path,
    level: str = "basic",
    state_path: Path | None = None,
    config: Mapping[str, object] | None = None,
) -> dict[str, object]:
    input_path = Path(input_path)
    state_path = Path(state_path) if state_path else None
    config = dict(config or {})
    signature = _signature(input_path, level, config)

    if state_path and state_path.exists():
        state = json.loads(state_path.read_text(encoding="utf-8"))
        if state.get("signature") == signature:
            return {
                "reused": True,
                "signature": signature,
                "issues": [Issue(**item) for item in state.get("issues", [])],
            }

    issues = check_syllabus(read_syllabus(input_path), level=level)
    result = {"reused": False, "signature": signature, "issues": issues}
    if state_path:
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "signature": signature,
                    "input_path": str(input_path),
                    "level": level,
                    "checked_at_utc": datetime.now(timezone.utc).isoformat(),
                    "issues": [asdict(issue) for issue in issues],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
    return result


def write_report(path: Path, input_path: Path, level: str, result: Mapping[str, object]) -> None:
    issues = result.get("issues", [])
    blockers = [issue for issue in issues if isinstance(issue, Issue) and issue.severity == "blocking"]
    warnings = [issue for issue in issues if isinstance(issue, Issue) and issue.severity == "warning"]
    lines = [
        "# 课程大纲审查报告",
        "",
        f"- 输入：`{Path(input_path).name}`",
        f"- 检查级别：`{level}`",
        f"- 构建签名：`{result.get('signature', '')}`",
        f"- 复用既有结果：`{str(bool(result.get('reused'))).lower()}`",
        f"- 阻断项：{len(blockers)}",
        f"- 警告项：{len(warnings)}",
        "",
    ]
    for title, severity in (("阻断项", "blocking"), ("警告项", "warning")):
        selected = [issue for issue in issues if isinstance(issue, Issue) and issue.severity == severity]
        lines.extend([f"## {title}", ""])
        if not selected:
            lines.append("无。")
        else:
            for issue in selected:
                lines.append(f"- **{issue.code}** `{issue.location}` — {issue.message}")
        lines.append("")
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def _assessment_sync_issues(
    syllabus: Mapping[str, object], assessment: Mapping[str, object]
) -> list[Issue]:
    issues: list[Issue] = []
    syllabus_meta = syllabus.get("metadata", {})
    assessment_meta = assessment.get("metadata", {})
    if not isinstance(syllabus_meta, dict) or not isinstance(assessment_meta, dict):
        return [
            _issue(
                "ASSESSMENT_SYNC_INVALID",
                "blocking",
                "Syllabus or assessment metadata is invalid",
                "控制区",
            )
        ]
    if str(syllabus_meta.get("assessment_version", "")) != str(
        assessment_meta.get("assessment_version", "")
    ):
        issues.append(
            _issue(
                "ASSESSMENT_VERSION_MISMATCH",
                "blocking",
                "Syllabus assessment version does not match assessment-plan.md",
                "控制区/assessment_version",
            )
        )
    if str(syllabus_meta.get("assessment_hash", "")) != str(
        assessment_meta.get("content_hash", "")
    ):
        issues.append(
            _issue(
                "ASSESSMENT_HASH_MISMATCH",
                "blocking",
                "Syllabus assessment hash does not match assessment-plan.md",
                "控制区/assessment_hash",
            )
        )

    syllabus_items = {
        row.get("assessment_id", ""): row for row in _table(syllabus, "课程评价")
    }
    source_items = {
        row.get("assessment_id", ""): row for row in _table(assessment, "考核项目")
    }
    if set(syllabus_items) != set(source_items):
        issues.append(
            _issue(
                "ASSESSMENT_ITEMS_MISMATCH",
                "blocking",
                "Materialized syllabus assessment IDs differ from assessment-plan.md",
                "课程评价",
            )
        )
    for assessment_id in sorted(set(syllabus_items) & set(source_items)):
        target = syllabus_items[assessment_id]
        source = source_items[assessment_id]
        for target_field, source_field in (("考核项目", "考核项目"), ("方式", "方式"), ("权重", "权重")):
            if target.get(target_field, "").strip() != source.get(source_field, "").strip():
                issues.append(
                    _issue(
                        "ASSESSMENT_CONTENT_MISMATCH",
                        "blocking",
                        f"Assessment {assessment_id} field {target_field} differs from assessment-plan.md",
                        f"课程评价/{assessment_id}",
                    )
                )
    return issues


def _set_run_font(run, east_asia: str, latin: str, size: float, bold: bool | None = None) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _set_cell_text(cell, text: str, *, bold: bool = False, centered: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "Syllabus Table Text"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    _set_run_font(run, "宋体", "Times New Roman", 9, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def _keep_row_together(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trpr.append(cant_split)


def _add_table(document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True, centered=True)
    _repeat_table_header(table.rows[0])
    _keep_row_together(table.rows[0])
    for values in rows:
        row = table.add_row()
        _keep_row_together(row)
        for index, value in enumerate(values):
            _set_cell_text(row.cells[index], value, centered=index == 0)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph()
    return table


def _add_info_table(document, data: Mapping[str, object]) -> None:
    info_rows = _table(data, "课程基本信息")
    info = {row.get("字段", ""): row.get("内容", "") for row in info_rows}
    table = document.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    name_row = table.add_row()
    _set_cell_text(name_row.cells[0], "课程名称（中英文）", bold=True, centered=True)
    name_value = info.get("课程中文名称", "")
    if info.get("课程英文名称", ""):
        name_value = f"{name_value}  {info['课程英文名称']}"
    merged = name_row.cells[1].merge(name_row.cells[3])
    _set_cell_text(merged, name_value, centered=True)

    pairs = [
        ("课程代码", "课程类别"),
        ("课程性质", "授课语言"),
        ("开课学期", "学分"),
        ("适用专业", "开课单位"),
        ("先修课程", "后续课程"),
    ]
    for left, right in pairs:
        if not info.get(left, "") and not info.get(right, ""):
            continue
        row = table.add_row()
        _set_cell_text(row.cells[0], left, bold=True, centered=True)
        _set_cell_text(row.cells[1], info.get(left, ""), centered=True)
        _set_cell_text(row.cells[2], right, bold=True, centered=True)
        _set_cell_text(row.cells[3], info.get(right, ""), centered=True)

    hour_fields = ["总学时", "理论学时", "实践学时", "实验学时", "上机学时", "项目学时"]
    hours = "；".join(f"{field}：{info.get(field, '0')}" for field in hour_fields)
    hour_row = table.add_row()
    _set_cell_text(hour_row.cells[0], "课程学时及分配", bold=True, centered=True)
    hour_cell = hour_row.cells[1].merge(hour_row.cells[3])
    _set_cell_text(hour_cell, hours, centered=True)

    intro = _section_prose(str(data.get("text", "")), "课程简介")
    intro_row = table.add_row()
    _set_cell_text(intro_row.cells[0], "课程简介", bold=True, centered=True)
    intro_cell = intro_row.cells[1].merge(intro_row.cells[3])
    _set_cell_text(intro_cell, intro)

    if info.get("大纲更新时间", ""):
        update_row = table.add_row()
        _set_cell_text(update_row.cells[0], "大纲更新时间", bold=True, centered=True)
        update_cell = update_row.cells[1].merge(update_row.cells[3])
        _set_cell_text(update_cell, info["大纲更新时间"], centered=True)
    for row in table.rows:
        _keep_row_together(row)
    document.add_paragraph()


def _add_heading(document, number: int, title: str) -> None:
    numeral = CHINESE_NUMERALS[number - 1] if number <= len(CHINESE_NUMERALS) else str(number)
    document.add_paragraph(f"{numeral}、{title}", style="Heading 1")


def _unit_rows(units: list[dict[str, str]]) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in units:
        requirements = "\n".join(
            part
            for part in (
                f"思政融入点：{row.get('思政融入点', '')}" if row.get("思政融入点", "") else "",
                f"预期学习成果：{row.get('预期学习成果', '')}" if row.get("预期学习成果", "") else "",
                f"教学方式：{row.get('教学方式', '')}" if row.get("教学方式", "") else "",
            )
            if part
        )
        rows.append(
            [
                f"{row.get('unit_id', '')} {row.get('单元名称', '')}".strip(),
                row.get("教学内容", ""),
                row.get("学时", ""),
                requirements,
                row.get("对应目标", ""),
            ]
        )
    return rows


def _render_document(
    syllabus: Mapping[str, object], assessment: Mapping[str, object], template_path: Path
):
    document = Document(template_path)
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    info = {row.get("字段", ""): row.get("内容", "") for row in _table(syllabus, "课程基本信息")}
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(f"《{info.get('课程中文名称', '')}》课程教学大纲")
    _set_run_font(title_run, "黑体", "Times New Roman", 16, True)

    section_number = 1
    _add_heading(document, section_number, "课程基本信息")
    _add_info_table(document, syllabus)
    section_number += 1

    _add_heading(document, section_number, "课程目标")
    goal_rows = [
        [
            f"{row.get('goal_id', '')} {row.get('类型', '')}".strip(),
            row.get("目标内容", ""),
            row.get("达成方法", ""),
        ]
        for row in _table(syllabus, "课程目标")
    ]
    _add_table(
        document,
        ["课程总目标", "课程子目标（参考学校定位、培养目标、毕业要求）", "达成方法"],
        goal_rows,
        [2.2, 9.5, 4.8],
    )
    section_number += 1

    units = _table(syllabus, "教学单元")
    theory = [row for row in units if any(word in row.get("类型", "") for word in ("理论", "混合"))]
    practice = [
        row
        for row in units
        if any(word in row.get("类型", "") for word in ("实践", "实验", "实训", "上机", "项目", "混合"))
    ]
    theory_hours = _integer(info.get("理论学时", "0")) or 0
    practice_hours = sum((_integer(info.get(field, "0")) or 0) for field in ("实践学时", "实验学时", "上机学时", "项目学时"))
    headers = ["单元/项目", "教学内容", "学时", "课程要求", "课程目标"]
    widths = [2.5, 3.8, 1.2, 7.0, 2.0]
    if theory_hours > 0:
        _add_heading(document, section_number, "理论教学环节及基本要求")
        _add_table(document, headers, _unit_rows(theory or units), widths)
        section_number += 1
    if practice_hours > 0:
        _add_heading(document, section_number, "实践（实验）教学环节及基本要求")
        _add_table(document, headers, _unit_rows(practice or units), widths)
        section_number += 1

    _add_heading(document, section_number, "课程评价")
    section_number += 1
    document.add_paragraph("（一）考核内容、考核方式与课程目标对应关系", style="Heading 2")
    coverage_rows = _table(assessment, "目标与单元覆盖")
    coverage_by_id: dict[str, dict[str, set[str]]] = {}
    for row in coverage_rows:
        item = coverage_by_id.setdefault(row.get("assessment_id", ""), {"goals": set(), "units": set()})
        item["goals"].add(row.get("goal_id", ""))
        item["units"].add(row.get("unit_id", ""))
    assessment_rows = []
    source_assessments = _table(assessment, "考核项目")
    for row in source_assessments:
        coverage = coverage_by_id.get(row.get("assessment_id", ""), {"goals": set(), "units": set()})
        assessment_rows.append(
            [
                row.get("assessment_id", ""),
                row.get("考核项目", ""),
                row.get("方式", ""),
                row.get("权重", ""),
                "、".join(sorted(coverage["units"])),
                "、".join(sorted(coverage["goals"])),
            ]
        )
    _add_table(
        document,
        ["编号", "考核项目", "方式", "成绩占比（%）", "覆盖单元", "课程目标"],
        assessment_rows,
        [1.4, 3.0, 2.2, 2.2, 3.8, 3.8],
    )
    document.add_paragraph("（二）考核方式与评分标准", style="Heading 2")
    criteria = _table(assessment, "评分标准")
    for item in source_assessments:
        assessment_id = item.get("assessment_id", "")
        document.add_paragraph(
            f"{assessment_id} {item.get('考核项目', '')}评分标准", style="Heading 2"
        )
        rubric_rows = [
            [
                row.get("评分指标", ""),
                row.get("指标权重", ""),
                row.get("优秀（90–100）", ""),
                row.get("良好（80–89）", ""),
                row.get("中等（70–79）", ""),
                row.get("及格（60–69）", ""),
                row.get("不及格（0–59）", ""),
            ]
            for row in criteria
            if row.get("assessment_id", "") == assessment_id
        ]
        _add_table(
            document,
            ["评分指标", "指标权重（%）", "优秀", "良好", "中等", "及格", "不及格"],
            rubric_rows,
            [2.2, 1.5, 2.6, 2.6, 2.6, 2.6, 2.6],
        )

    _add_heading(document, section_number, "参考书目及学习资料")
    resources = _table(syllabus, "教材与学习资料")
    _add_table(
        document,
        ["类型", "资料信息"],
        [[row.get("类型", ""), row.get("资料信息", "")] for row in resources],
        [3.0, 13.5],
    )

    signing = _table(syllabus, "签批信息")
    if signing:
        _add_table(
            document,
            ["角色", "姓名", "日期"],
            [[row.get("角色", ""), row.get("姓名", ""), row.get("日期", "")] for row in signing],
            [3.0, 7.0, 6.5],
        )

    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""
    properties.keywords = ""
    return document


def inspect_docx(path: Path) -> InspectionResult:
    path = Path(path)
    findings: list[str] = []
    visible_parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        story_names = sorted(
            name
            for name in names
            if name == "word/document.xml"
            or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
            or name in {"word/footnotes.xml", "word/endnotes.xml"}
        )
        story_xml = []
        for name in story_names:
            payload = archive.read(name)
            story_xml.append(payload.decode("utf-8", errors="replace"))
            root = ElementTree.fromstring(payload)
            for element in root.iter():
                if element.tag == qn("w:t") and element.text:
                    visible_parts.append(element.text)
        joined_xml = "\n".join(story_xml)
        visible_text = "\n".join(visible_parts)

        if any(name.startswith("word/comments") for name in names) or "<w:commentReference" in joined_xml:
            findings.append("COMMENTS_PRESENT")
        if re.search(r"<w:(?:ins|del|moveFrom|moveTo)(?:\s|>)", joined_xml):
            findings.append("TRACKED_CHANGES_PRESENT")
        if "<w:vanish" in joined_xml or "<w:webHidden" in joined_xml:
            findings.append("HIDDEN_TEXT_PRESENT")
        if "[待填写]" in visible_text or "to_confirm" in visible_text:
            findings.append("UNRESOLVED_PLACEHOLDER")
        if any(pattern.search(visible_text) for pattern in AI_AUTHORSHIP_PATTERNS):
            findings.append("AI_AUTHORSHIP_MARKER")

        metadata = ""
        for name in ("docProps/core.xml", "docProps/custom.xml", "docProps/app.xml"):
            if name in names:
                metadata += archive.read(name).decode("utf-8", errors="replace")
        if re.search(r"OpenAI|ChatGPT|Codex|GPT-[0-9]", metadata, re.I):
            findings.append("TOOL_METADATA_PRESENT")

    codes = tuple(dict.fromkeys(findings))
    return InspectionResult(
        clean=not codes,
        codes=codes,
        findings=codes,
        visible_text=visible_text,
    )


def _render_signature(
    input_path: Path,
    assessment_path: Path,
    template_path: Path,
    config: Mapping[str, object],
) -> str:
    digest = hashlib.sha256()
    for path in (input_path, assessment_path, template_path):
        digest.update(Path(path).read_bytes())
    digest.update(SCRIPT_VERSION.encode("ascii"))
    digest.update(json.dumps(config, ensure_ascii=False, sort_keys=True).encode("utf-8"))
    return digest.hexdigest()


def render_syllabus(
    input_path: Path,
    assessment_path: Path,
    template_path: Path,
    output_path: Path,
    *,
    state_path: Path | None = None,
    config: Mapping[str, object] | None = None,
) -> dict[str, object]:
    input_path = Path(input_path)
    assessment_path = Path(assessment_path)
    template_path = Path(template_path)
    output_path = Path(output_path)
    state_path = Path(state_path) if state_path else None
    config = dict(config or {})

    syllabus = read_syllabus(input_path)
    blockers = [issue for issue in check_syllabus(syllabus, "basic") if issue.severity == "blocking"]
    if blockers:
        raise ValueError("blocking syllabus issues: " + ", ".join(issue.code for issue in blockers))
    assessment = read_assessment(assessment_path)
    sync_issues = _assessment_sync_issues(syllabus, assessment)
    if sync_issues:
        raise ValueError("blocking assessment sync issues: " + ", ".join(issue.code for issue in sync_issues))
    if not template_path.exists():
        raise ValueError(f"Word template does not exist: {template_path}")

    signature = _render_signature(input_path, assessment_path, template_path, config)
    if state_path and state_path.exists() and output_path.exists():
        state = json.loads(state_path.read_text(encoding="utf-8"))
        if state.get("signature") == signature:
            inspection = inspect_docx(output_path)
            if inspection.clean:
                return {"reused": True, "signature": signature, "inspection": inspection}

    document = _render_document(syllabus, assessment, template_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    inspection = inspect_docx(output_path)
    if not inspection.clean:
        raise ValueError("generated Word failed inspection: " + ", ".join(inspection.codes))

    if state_path:
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "signature": signature,
                    "input_path": str(input_path),
                    "assessment_path": str(assessment_path),
                    "template_path": str(template_path),
                    "output_path": str(output_path),
                    "rendered_at_utc": datetime.now(timezone.utc).isoformat(),
                    "inspection_codes": list(inspection.codes),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
    return {"reused": False, "signature": signature, "inspection": inspection}


def write_inspection_report(path: Path, input_path: Path, result: InspectionResult) -> None:
    lines = [
        "# Word 终稿检查报告",
        "",
        f"- 输入：`{Path(input_path).name}`",
        f"- 结果：{'通过' if result.clean else '未通过'}",
        f"- 问题数：{len(result.codes)}",
        "",
        "## 检查结果",
        "",
    ]
    if result.clean:
        lines.append("无批注、修订、隐藏文字、未解决占位符或生成工具署名。")
    else:
        lines.extend(f"- `{code}`" for code in result.codes)
    Path(path).write_text("\n".join(lines) + "\n", encoding="utf-8")


def _load_config(path: Path | None) -> dict[str, object]:
    if path is None:
        return {}
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Config must be a JSON object")
    return payload


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Check and compile course syllabus artifacts")
    parser.add_argument("--version", action="version", version=SCRIPT_VERSION)
    subparsers = parser.add_subparsers(dest="command", required=True)
    check = subparsers.add_parser("check", help="Check a syllabus working Markdown file")
    check.add_argument("--input", required=True, type=Path)
    check.add_argument("--report", required=True, type=Path)
    check.add_argument("--level", choices=sorted(VALID_LEVELS), default=None)
    check.add_argument("--config", type=Path)
    check.add_argument("--state", type=Path)
    render = subparsers.add_parser("render", help="Compile a syllabus Markdown file to Word")
    render.add_argument("--input", required=True, type=Path)
    render.add_argument("--assessment", required=True, type=Path)
    render.add_argument("--template", required=True, type=Path)
    render.add_argument("--output", required=True, type=Path)
    render.add_argument("--config", type=Path)
    render.add_argument("--state", type=Path)
    inspect = subparsers.add_parser("inspect", help="Inspect a final syllabus Word document")
    inspect.add_argument("--input", required=True, type=Path)
    inspect.add_argument("--report", required=True, type=Path)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = _build_parser()
    args = parser.parse_args(argv)
    if args.command == "check":
        try:
            config = _load_config(args.config)
            level = args.level or str(config.get("review_level", "basic"))
            result = check_file(args.input, level=level, state_path=args.state, config=config)
            write_report(args.report, args.input, level, result)
            issues = result["issues"]
            return 2 if any(issue.severity == "blocking" for issue in issues) else 0
        except (OSError, ValueError, yaml.YAMLError, json.JSONDecodeError) as exc:
            parser.error(str(exc))
    if args.command == "render":
        try:
            config = _load_config(args.config)
            render_syllabus(
                args.input,
                args.assessment,
                args.template,
                args.output,
                state_path=args.state,
                config=config,
            )
            return 0
        except (OSError, ValueError, yaml.YAMLError, json.JSONDecodeError, zipfile.BadZipFile) as exc:
            parser.error(str(exc))
    if args.command == "inspect":
        try:
            result = inspect_docx(args.input)
            write_inspection_report(args.report, args.input, result)
            return 0 if result.clean else 2
        except (OSError, ValueError, zipfile.BadZipFile) as exc:
            parser.error(str(exc))
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
